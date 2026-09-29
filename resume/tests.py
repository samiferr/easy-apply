import io
import json
import shutil
import tempfile
from datetime import date
from unittest.mock import patch

import docx
from django.conf import settings
from django.contrib.auth import get_user_model
from django.contrib.messages import get_messages
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from accounts.models import Profile
from core.models import AITask
from core.testing import fake_deepseek
from education.models import Certificate, Degree
from experience.models import ExperienceHighlight, WorkExperience
from jobs.models import JobElement, JobPost, JobSection
from languages.models import Language, UserLanguage
from skills.models import SkillCategory, UserSkill
from staffportal.models import Plan, UsageMetric, UsageRecord
from staffportal.services import runtime_settings

from .models import ResumeImport, TailoredResume
from .services.pdf import markdown_to_flowables, render_markdown_pdf
from .services.tailored import (
    EDUCATION,
    EXPERIENCE,
    LANGUAGES,
    SKILLS,
    SUMMARY,
    build_job_payload,
    generate_tailored_resume,
    load_template_sections,
    render_markdown,
)

User = get_user_model()

# Tests always run with DEBUG=False, where the manifest static-files storage
# insists on a `collectstatic` having been run; plain storage keeps the
# template-rendering tests below independent of the build step.
TEST_STORAGES = {
    **settings.STORAGES,
    "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"},
}

AI_RESPONSE = {
    "professional_summary": "Backend engineer with five years building Django services.",
    "skills": [
        {"group": "Programming", "items": ["Python", "Django"]},
        {"group": "Empty group", "items": []},
    ],
    "experience": [
        {
            "job_title": "Backend Engineer",
            "company": "Acme Corp",
            "location": "Montreal, QC",
            "dates": "Jan 2020 – Present",
            "highlights": ["Built a payments API serving 2M requests a day."],
        }
    ],
    "education": [
        {
            "title": "BSc Computer Science",
            "organization": "McGill University",
            "dates": "2015 – 2019",
            "details": "",
        }
    ],
    "languages": ["English — Native", "French — Fluent"],
    "_model": "deepseek-test",
}


class TailoredResumeTestMixin:
    LANGUAGE = "en"

    def setUp(self):
        self.user = User.objects.create_user(
            email="jane@example.com", password="pw-for-tests-123", first_name="Jane", last_name="Doe"
        )
        self.profile = Profile.objects.get(user=self.user)
        self.profile.name = "Main"
        self.profile.language = self.LANGUAGE
        self.profile.phone = "+1 555 0100"
        self.profile.location = "Montreal, QC"
        self.profile.save()
        self.job = JobPost.objects.create(
            profile=self.profile,
            status=JobPost.STATUS_COMPLETED,
            title="Senior Backend Engineer",
            company_name="Globex",
        )
        section = JobSection.objects.create(
            job=self.job, key="required_technical_skills"
        )
        JobElement.objects.create(
            section=section,
            text="5+ years of Python",
            match_status=JobElement.STRONG,
            match_evidence="5 years as Backend Engineer at Acme Corp.",
        )

    def add_profile_content(self):
        skill_category, _ = SkillCategory.objects.get_or_create(
            name="Programming Languages", kind=SkillCategory.TECHNICAL
        )
        UserSkill.objects.create(
            profile=self.profile, category=skill_category, name="Python", level=UserSkill.EXPERT
        )
        experience = WorkExperience.objects.create(
            profile=self.profile,
            job_title="Backend Engineer",
            company="Acme Corp",
            location="Montreal, QC",
            start_date=date(2020, 1, 1),
            is_current=True,
        )
        ExperienceHighlight.objects.create(experience=experience, text="Built a payments API.")


class TemplateSectionTests(TestCase):
    def test_sections_come_from_the_template_file(self):
        self.assertEqual(
            load_template_sections(),
            [
                (SUMMARY, "PROFESSIONAL SUMMARY"),
                (SKILLS, "SKILLS"),
                (EXPERIENCE, "WORKING EXPERIENCE"),
                (EDUCATION, "EDUCATION & PROFESSIONAL DEVELOPMENT"),
                (LANGUAGES, "LANGUAGES"),
            ],
        )


class RenderMarkdownTests(TailoredResumeTestMixin, TestCase):
    def test_header_and_sections_follow_the_template(self):
        markdown = render_markdown(self.profile, AI_RESPONSE)

        self.assertTrue(markdown.startswith("**Jane Doe**"))
        self.assertIn("Montreal, QC  |  +1 555 0100  |  jane@example.com", markdown)

        headings = [line for line in markdown.splitlines() if line.startswith("## ")]
        self.assertEqual(
            headings,
            [
                "## **PROFESSIONAL SUMMARY**",
                "## **SKILLS**",
                "## **WORKING EXPERIENCE**",
                "## **EDUCATION & PROFESSIONAL DEVELOPMENT**",
                "## **LANGUAGES**",
            ],
        )
        self.assertIn("- **Programming:** Python, Django", markdown)
        self.assertNotIn("Empty group", markdown)
        self.assertIn("### **Backend Engineer** — Acme Corp", markdown)
        self.assertIn("*Jan 2020 – Present · Montreal, QC*", markdown)
        self.assertIn("- Built a payments API serving 2M requests a day.", markdown)
        self.assertIn("- English — Native", markdown)

    def test_sections_without_content_are_dropped(self):
        markdown = render_markdown(self.profile, {"professional_summary": "Just a summary."})

        self.assertIn("## **PROFESSIONAL SUMMARY**", markdown)
        self.assertNotIn("## **SKILLS**", markdown)
        self.assertNotIn("## **LANGUAGES**", markdown)

    def test_malformed_ai_content_is_ignored_rather_than_rendered(self):
        markdown = render_markdown(
            self.profile,
            {
                "professional_summary": 42,
                "skills": "not a list",
                "experience": [None, {"highlights": ["orphan"]}],
                "languages": ["", None, "English — Native"],
            },
        )

        self.assertNotIn("42", markdown)
        self.assertNotIn("orphan", markdown)
        self.assertIn("- English — Native", markdown)


class BuildJobPayloadTests(TailoredResumeTestMixin, TestCase):
    def test_requirements_carry_their_profile_match(self):
        payload = build_job_payload(self.job)

        self.assertEqual(payload["title"], "Senior Backend Engineer")
        self.assertEqual(payload["company"], "Globex")
        self.assertEqual(len(payload["requirements"]), 1)
        requirement = payload["requirements"][0]
        self.assertEqual(requirement["text"], "5+ years of Python")
        self.assertEqual(requirement["profile_match"], "strong")
        # Categories now come from the fixed section enum, not free-form AI names.
        self.assertEqual(requirement["category"], "Required Technical Skills")


class GenerateTailoredResumeTests(TailoredResumeTestMixin, TestCase):
    def test_empty_profile_is_skipped_before_calling_the_ai(self):
        with patch("resume.services.tailored.call_deepseek_json") as call:
            result = generate_tailored_resume(self.job)

        call.assert_not_called()
        self.assertEqual(result, {"skipped": "empty_profile"})
        self.assertFalse(TailoredResume.objects.exists())

    def test_draft_is_stored_as_markdown(self):
        self.add_profile_content()
        with patch(
            "resume.services.tailored.call_deepseek_json", return_value=dict(AI_RESPONSE)
        ) as call:
            result = generate_tailored_resume(self.job)

        call.assert_called_once()
        tailored_resume = result["tailored_resume"]
        self.assertEqual(tailored_resume.job, self.job)
        self.assertEqual(tailored_resume.ai_model, "deepseek-test")
        self.assertFalse(tailored_resume.edited_by_user)
        self.assertIn("## **PROFESSIONAL SUMMARY**", tailored_resume.markdown)
        self.assertIsNotNone(tailored_resume.generated_at)

    def test_regenerating_replaces_the_existing_draft(self):
        self.add_profile_content()
        TailoredResume.objects.create(
            profile=self.profile, job=self.job, markdown="old draft", edited_by_user=True
        )

        with patch("resume.services.tailored.call_deepseek_json", return_value=dict(AI_RESPONSE)):
            generate_tailored_resume(self.job)

        self.assertEqual(TailoredResume.objects.count(), 1)
        tailored_resume = TailoredResume.objects.get()
        self.assertNotIn("old draft", tailored_resume.markdown)
        self.assertFalse(tailored_resume.edited_by_user)


class MarkdownPDFTests(TestCase):
    MARKDOWN = (
        "**Jane Doe**\n\n"
        "Montreal, QC  |  jane@example.com\n\n"
        "## **PROFESSIONAL SUMMARY**\n\n"
        "Backend engineer & problem solver.\n\n"
        "## **WORKING EXPERIENCE**\n\n"
        "### **Backend Engineer** — Acme Corp\n"
        "*Jan 2020 – Present*\n\n"
        "- Built a payments API.\n"
        "- See [the docs](https://example.com/docs).\n"
    )

    def test_renders_a_pdf_document(self):
        pdf_bytes = render_markdown_pdf(self.MARKDOWN, title="Resume", author="Jane Doe")

        self.assertTrue(pdf_bytes.startswith(b"%PDF"))
        self.assertGreater(len(pdf_bytes), 1000)

    def test_pdf_contains_the_resume_text(self):
        import io

        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(render_markdown_pdf(self.MARKDOWN)))
        text = "\n".join(page.extract_text() for page in reader.pages)

        self.assertIn("Jane Doe", text)
        self.assertIn("PROFESSIONAL SUMMARY", text)
        self.assertIn("Backend engineer & problem solver.", text)
        self.assertIn("Built a payments API.", text)

    def test_markup_characters_do_not_leak_into_the_page(self):
        story = markdown_to_flowables("## **SKILLS**\n\n- **Python** and *Django*\n")
        rendered = " ".join(getattr(flowable, "text", "") for flowable in story)

        self.assertIn("SKILLS", rendered)
        self.assertIn("<b>Python</b>", rendered)
        self.assertIn("<i>Django</i>", rendered)
        self.assertNotIn("**", rendered)

    def test_empty_markdown_still_produces_a_pdf(self):
        self.assertTrue(render_markdown_pdf("   ").startswith(b"%PDF"))

    def test_unbalanced_markup_falls_back_to_plain_text(self):
        # `*one **two* three**` converts to overlapping tags; the export
        # must still succeed rather than 500 on a hand-edited line.
        with self.assertLogs("resume.services.pdf", level="WARNING"):
            pdf_bytes = render_markdown_pdf("*one **two* three**\n\n## SKILLS\n")

        self.assertTrue(pdf_bytes.startswith(b"%PDF"))

    def test_html_in_the_markdown_is_escaped(self):
        story = markdown_to_flowables("Sales <b>up</b> & to the right\n")
        rendered = " ".join(getattr(flowable, "text", "") for flowable in story)

        self.assertIn("&amp;", rendered)
        self.assertNotIn("<b>up</b>", rendered)


@override_settings(STORAGES=TEST_STORAGES)
class TailoredResumeViewTests(TailoredResumeTestMixin, TestCase):
    def setUp(self):
        super().setUp()
        self.add_profile_content()
        self.client.force_login(self.user)

    def generate(self):
        with patch("resume.services.tailored.call_deepseek_json", return_value=dict(AI_RESPONSE)):
            return self.client.post(reverse("resume:tailored_generate", args=[self.job.pk]))

    def test_generate_redirects_to_the_editor(self):
        response = self.generate()

        self.assertRedirects(response, reverse("resume:tailored", args=[self.job.pk]))
        self.assertTrue(TailoredResume.objects.filter(job=self.job).exists())

    def test_generate_requires_a_completed_analysis(self):
        self.job.status = JobPost.STATUS_FAILED
        self.job.save(update_fields=["status"])

        response = self.client.post(reverse("resume:tailored_generate", args=[self.job.pk]))

        self.assertRedirects(response, self.job.get_absolute_url())
        self.assertFalse(TailoredResume.objects.exists())

    def test_editing_marks_the_draft_as_user_edited(self):
        self.generate()

        response = self.client.post(
            reverse("resume:tailored", args=[self.job.pk]),
            {"markdown": "**Jane Doe**\n\n## **SKILLS**\n\n- Python\n", "action": "save"},
        )

        self.assertRedirects(response, reverse("resume:tailored", args=[self.job.pk]))
        tailored_resume = TailoredResume.objects.get(job=self.job)
        self.assertTrue(tailored_resume.edited_by_user)
        self.assertIn("- Python", tailored_resume.markdown)

    def test_save_and_export_returns_the_pdf_of_the_edited_text(self):
        self.generate()

        response = self.client.post(
            reverse("resume:tailored", args=[self.job.pk]),
            {"markdown": "**Jane Doe**\n\n## **SKILLS**\n\n- Rust\n", "action": "pdf"},
        )

        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertIn("attachment; filename=", response["Content-Disposition"])
        self.assertTrue(response.content.startswith(b"%PDF"))
        self.assertIn("- Rust", TailoredResume.objects.get(job=self.job).markdown)

    def test_empty_markdown_is_rejected(self):
        self.generate()

        response = self.client.post(
            reverse("resume:tailored", args=[self.job.pk]), {"markdown": "   ", "action": "save"}
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "can&#x27;t be empty")

    def test_pdf_and_markdown_downloads(self):
        self.generate()

        pdf = self.client.get(reverse("resume:tailored_pdf", args=[self.job.pk]))
        markdown = self.client.get(reverse("resume:tailored_markdown", args=[self.job.pk]))

        self.assertTrue(pdf.content.startswith(b"%PDF"))
        self.assertEqual(pdf["Content-Disposition"], 'attachment; filename="jane-doe-globex-resume.pdf"')
        self.assertEqual(markdown["Content-Type"], "text/markdown; charset=utf-8")
        self.assertIn(b"## **PROFESSIONAL SUMMARY**", markdown.content)

    def test_another_users_resume_is_not_reachable(self):
        self.generate()
        other = User.objects.create_user(email="mallory@example.com", password="pw-for-tests-123")
        self.client.force_login(other)

        for name in ("resume:tailored", "resume:tailored_pdf", "resume:tailored_markdown"):
            self.assertEqual(self.client.get(reverse(name, args=[self.job.pk])).status_code, 404)
        self.assertEqual(
            self.client.post(reverse("resume:tailored_generate", args=[self.job.pk])).status_code, 404
        )

    def test_login_is_required(self):
        self.client.logout()

        response = self.client.get(reverse("resume:tailored", args=[self.job.pk]))

        self.assertEqual(response.status_code, 302)
        self.assertIn(reverse("accounts:login"), response["Location"])


class FrenchProfileResumeTests(TailoredResumeTestMixin, TestCase):
    """A French profile produces a French resume, whatever the UI is set to."""

    LANGUAGE = "fr"

    def test_the_prompt_asks_for_french_and_the_headings_are_french(self):
        from django.utils import translation

        self.add_profile_content()
        captured = {}

        def fake_call(system_prompt, user_content, **kwargs):
            captured["payload"] = user_content
            return dict(AI_RESPONSE)

        with translation.override("en"), patch(
            "resume.services.tailored.call_deepseek_json", side_effect=fake_call
        ):
            result = generate_tailored_resume(self.job)

        self.assertIn("in French", captured["payload"])
        markdown = result["tailored_resume"].markdown
        self.assertIn("## **RÉSUMÉ PROFESSIONNEL**", markdown)
        self.assertNotIn("## **PROFESSIONAL SUMMARY**", markdown)

    def test_the_profile_snapshot_reaches_the_model_in_french(self):
        self.add_profile_content()
        captured = {}

        def fake_call(system_prompt, user_content, **kwargs):
            captured["payload"] = user_content
            return dict(AI_RESPONSE)

        with patch("resume.services.tailored.call_deepseek_json", side_effect=fake_call):
            generate_tailored_resume(self.job)

        # The current role's end marker is the giveaway: the snapshot carries
        # display strings, and they must already be French when the model sees
        # them.
        self.assertIn("Aujourd’hui", captured["payload"])
        self.assertNotIn("Jan 2020 – Present", captured["payload"])

    def test_the_resume_is_parsed_in_the_profile_language(self):
        from resume.models import ResumeImport
        from resume.services import importer

        upload = ResumeImport.objects.create(profile=self.profile, file="resumes/x.txt")
        captured = {}

        def fake_analyze(raw_text, soft, technical, language="en"):
            captured["language"] = language
            return {"profile": {}}

        with patch.object(importer, "extract_resume_text", return_value="CV text"), \
             patch.object(importer, "analyze_resume_text", side_effect=fake_analyze):
            importer.run_analysis(upload)

        self.assertEqual(captured["language"], "fr")


# ---------------------------------------------------------------------------
# Characterization tests — pin the use cases before they move into services.py.
#
# Real tasks, real extractor, real importer; only the HTTP call to the AI
# provider is stubbed, so none of this depends on where the code lives.
# ---------------------------------------------------------------------------
def flash(response):
    return [str(message) for message in get_messages(response.wsgi_request)]


RESUME_TEXT = (
    "Jane Doe. Backend engineer with five years of Python experience building "
    "payment systems at Acme Corp in Montreal."
)

RESUME_REPLY = {
    "profile": {
        "first_name": "Jane", "last_name": "Doe", "headline": "Backend engineer",
        "phone": "+1 555 0100", "location": "Montreal, QC", "bio": "Ships payment systems.",
        "linkedin_url": "https://linkedin.com/in/jane", "portfolio_url": "", "github_url": "",
    },
    "soft_skills": [
        {"name": "Mentoring", "category": "Leadership", "level": "advanced"},
    ],
    "technical_skills": [
        {"name": "Python", "category": "Programming Languages", "level": "expert"},
        {"name": "Zig", "category": "Brand New Category", "level": "wizard"},
        {"name": "Docker", "category": "", "level": ""},
    ],
    "languages": [
        {"name": "French", "proficiency": "native"},
        {"name": "Klingon", "proficiency": "sparkly"},
    ],
    "experience": [
        {
            "job_title": "Backend Engineer", "company": "Acme", "location": "Montreal",
            "employment_type": "full_time", "start_date": "2020-01-15", "end_date": None,
            "is_current": True, "highlights": ["Built the payments API.", "Cut latency by half."],
        },
        {
            "job_title": "Developer", "company": "Globex", "location": "",
            "employment_type": "made_up", "start_date": "2017", "end_date": "2019-12",
            "is_current": False, "highlights": [],
        },
        {"job_title": "", "company": "No title Inc", "highlights": []},
    ],
    "degrees": [
        {"school": "McGill", "degree": "BSc", "field_of_study": "CS", "start_date": "2013-09-01",
         "end_date": "2017-05-01", "is_current": False, "grade": "First"},
        {"school": "", "degree": "Orphan"},
    ],
    "certificates": [
        {"name": "AWS SAA", "issuing_organization": "Amazon", "issue_date": "2023-01-15",
         "expiry_date": "2026-01-15", "does_not_expire": False, "credential_id": "ABC",
         "credential_url": "https://aws.example/abc"},
        {"name": "No issuer", "issuing_organization": ""},
    ],
    "_model": "deepseek-test",
}


class TempMediaMixin:
    """Uploads go to a throwaway MEDIA_ROOT, never to the real one."""

    def setUp(self):
        super().setUp()
        self.media_root = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.media_root, ignore_errors=True)
        override = override_settings(MEDIA_ROOT=self.media_root, STORAGES=TEST_STORAGES)
        override.enable()
        self.addCleanup(override.disable)


def make_user_profile(email="cv@example.com", language="en", **user_fields):
    user = User.objects.create_user(email=email, password="pw-for-tests-123", **user_fields)
    profile = Profile.objects.get(user=user)
    profile.language = language
    profile.save(update_fields=["language"])
    return user, profile


class ExtractorTests(TestCase):
    """UC-04.2, step 3: plain text out of whatever the user uploaded."""

    def extract(self, name, content):
        from .services.extractor import extract_resume_text

        return extract_resume_text(SimpleUploadedFile(name, content))

    def assertRejected(self, name, content, message):
        from .services.extractor import ResumeExtractError

        with self.assertRaisesMessage(ResumeExtractError, message):
            self.extract(name, content)

    def docx_bytes(self, *paragraphs, table_cells=()):
        document = docx.Document()
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)
        if table_cells:
            table = document.add_table(rows=1, cols=len(table_cells))
            for cell, text in zip(table.rows[0].cells, table_cells):
                cell.text = text
        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    def test_a_text_file_is_decoded_and_stripped(self):
        self.assertEqual(self.extract("cv.TXT", f"  {RESUME_TEXT}\n".encode()), RESUME_TEXT)

    def test_a_text_file_with_bad_bytes_is_still_read(self):
        text = self.extract("cv.txt", RESUME_TEXT.encode() + b"\xff\xfe more")
        self.assertIn(RESUME_TEXT, text)

    def test_too_little_text_is_refused_for_every_format(self):
        self.assertRejected("cv.txt", b"too short", "That file doesn't seem to contain enough text")
        self.assertRejected(
            "cv.docx", self.docx_bytes("short"), "We couldn't find enough readable text"
        )

    def test_a_docx_yields_its_paragraphs_and_table_cells(self):
        text = self.extract(
            "cv.docx",
            self.docx_bytes(RESUME_TEXT, "", "Second paragraph.", table_cells=("Skill", "Python")),
        )
        self.assertIn(RESUME_TEXT, text)
        self.assertIn("Second paragraph.", text)
        self.assertTrue(text.endswith("Skill\nPython"))

    def test_a_pdf_yields_its_text(self):
        from .services.pdf import render_markdown_pdf

        pdf = render_markdown_pdf(f"**Jane Doe**\n\n## **SUMMARY**\n\n{RESUME_TEXT}\n")
        text = self.extract("cv.pdf", pdf)
        self.assertIn("Jane Doe", text)
        self.assertIn("Backend engineer", text)

    def test_broken_files_get_friendly_errors(self):
        self.assertRejected("cv.pdf", b"not a pdf at all", "Couldn't read that PDF")
        self.assertRejected("cv.docx", b"not a docx", "Couldn't read that DOCX file")

    def test_an_unsupported_type_is_refused(self):
        self.assertRejected(
            "cv.rtf", b"{\\rtf1 hello}", "Unsupported file type. Please upload a PDF, DOCX, or TXT file."
        )


@override_settings(STORAGES=TEST_STORAGES)
class ResumeUploadTests(TempMediaMixin, TestCase):
    """UC-04.1 and UC-04.2: upload, queue, extract, parse, store."""

    def setUp(self):
        super().setUp()
        runtime_settings.invalidate()
        self.addCleanup(runtime_settings.invalidate)
        Plan.objects.create(
            slug="free", name="Free", price_cents=0, is_default=True, monthly_resume_imports=1
        )
        self.user, self.profile = make_user_profile()
        self.client.force_login(self.user)

    def upload(self, name="cv.txt", content=RESUME_TEXT.encode()):
        return self.client.post(
            reverse("resume:upload"), {"file": SimpleUploadedFile(name, content)}
        )

    def imports_used(self):
        record = UsageRecord.objects.filter(
            user=self.user, metric=UsageMetric.RESUME_IMPORT
        ).first()
        return record.count if record else 0

    def test_an_upload_is_parsed_in_the_background_and_lands_on_the_review_page(self):
        with fake_deepseek(lambda request: dict(RESUME_REPLY)) as calls:
            response = self.upload()

        upload = ResumeImport.objects.get()
        self.assertRedirects(
            response, reverse("resume:review", args=[upload.pk]), fetch_redirect_response=False
        )
        self.assertEqual(upload.profile, self.profile)
        self.assertEqual(upload.original_filename, "cv.txt")
        self.assertTrue(upload.file.name.startswith(f"resumes/profile_{self.profile.pk}/cv"))
        self.assertEqual(upload.status, ResumeImport.STATUS_COMPLETED)
        self.assertEqual(upload.raw_text, RESUME_TEXT)
        self.assertEqual(upload.ai_model, "deepseek-test")
        self.assertEqual(upload.ai_response["profile"]["headline"], "Backend engineer")
        self.assertIsNotNone(upload.analyzed_at)
        task = AITask.latest_for(upload, AITask.RESUME_IMPORT)
        self.assertEqual((task.state, task.steps_total, task.current_step), (AITask.DONE, 2, "Resume ready to review"))
        self.assertEqual(self.imports_used(), 1)
        self.assertEqual(len(calls), 1)

    def test_the_model_is_asked_in_the_profile_language_about_the_existing_categories(self):
        _user, french = make_user_profile("cv-fr@example.com", language="fr")
        self.client.force_login(french.user)
        with fake_deepseek(lambda request: dict(RESUME_REPLY)) as calls:
            self.upload()
        (call,) = calls
        self.assertIn("expert resume parser", call.system)
        self.assertIn("in French", call.user)
        self.assertIn("--- Resume text ---\n" + RESUME_TEXT, call.user)
        for category in SkillCategory.objects.all():
            self.assertIn(category.name, call.user)
        self.assertEqual(call.temperature, 0.2)

    def test_the_resume_text_is_truncated_before_it_is_sent(self):
        with fake_deepseek(lambda request: dict(RESUME_REPLY)) as calls:
            self.upload(content=b"B" * 25000)
        self.assertEqual(calls[0].user.count("B"), 20000)

    def test_an_unreadable_file_fails_the_import_without_calling_the_model(self):
        with fake_deepseek(lambda request: dict(RESUME_REPLY)) as calls:
            self.upload(content=b"tiny")
        upload = ResumeImport.objects.get()
        self.assertEqual(calls, [])
        self.assertEqual(upload.status, ResumeImport.STATUS_FAILED)
        self.assertEqual(upload.error_message, "That file doesn't seem to contain enough text to analyze.")
        task = AITask.latest_for(upload, AITask.RESUME_IMPORT)
        self.assertEqual(task.state, AITask.FAILED)
        self.assertEqual(task.error_message, upload.error_message)

    def test_a_model_failure_is_recorded_on_the_import_and_the_task(self):
        with fake_deepseek(lambda request: 401):
            self.upload()
        upload = ResumeImport.objects.get()
        self.assertEqual(upload.status, ResumeImport.STATUS_FAILED)
        self.assertIn("rejected our API key", upload.error_message)
        self.assertEqual(AITask.latest_for(upload, AITask.RESUME_IMPORT).state, AITask.FAILED)

    def test_a_missing_api_key_is_reported(self):
        with override_settings(DEEPSEEK_API_KEY=""):
            self.upload()
        upload = ResumeImport.objects.get()
        self.assertEqual(upload.status, ResumeImport.STATUS_FAILED)
        self.assertIn("DEEPSEEK_API_KEY", upload.error_message)

    def test_an_exhausted_allowance_creates_nothing(self):
        runtime_settings.set_value("enforce_quotas", True)
        with fake_deepseek(lambda request: dict(RESUME_REPLY)):
            self.upload()
            response = self.upload()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            flash(response)[-1],
            "You've used all 1 of this month's resume imports. Your allowance resets at the "
            "start of your next billing period.",
        )
        self.assertEqual(ResumeImport.objects.count(), 1)
        self.assertEqual(self.imports_used(), 1)

    def test_the_kill_switch_blocks_an_upload(self):
        runtime_settings.set_value("ai_features_enabled", False)
        response = self.upload()
        self.assertEqual(response.status_code, 200)
        self.assertIn("temporarily unavailable", flash(response)[0])
        self.assertFalse(ResumeImport.objects.exists())
        self.assertEqual(self.imports_used(), 0)

    def test_the_form_refuses_the_wrong_type_and_an_oversized_file(self):
        response = self.upload("cv.exe", b"MZ" + b"x" * 100)
        self.assertEqual(response.status_code, 200)
        self.assertIn("file", response.context["form"].errors)
        with override_settings(RESUME_MAX_UPLOAD_BYTES=10):
            response = self.upload()
        self.assertIn("too large", " ".join(response.context["form"].errors["file"]))
        self.assertFalse(ResumeImport.objects.exists())
        self.assertEqual(self.imports_used(), 0)


@override_settings(STORAGES=TEST_STORAGES)
class ResumeReviewTests(TempMediaMixin, TestCase):
    """UC-04.3 and UC-04.4: what is new, what is already there, and applying the choice."""

    def setUp(self):
        super().setUp()
        self.user, self.profile = make_user_profile(first_name="", last_name="")
        self.client.force_login(self.user)
        self.upload = ResumeImport.objects.create(
            profile=self.profile, file="resumes/x.txt", original_filename="x.txt",
            status=ResumeImport.STATUS_COMPLETED, ai_response=json.loads(json.dumps(RESUME_REPLY)),
        )
        self.url = reverse("resume:review", args=[self.upload.pk])

    def sections(self):
        from .services.importer import build_review_sections

        return build_review_sections(self.upload, self.profile)

    # --- the checklist -----------------------------------------------------
    def test_profile_fields_are_only_offered_where_the_profile_is_empty(self):
        keys = [field["key"] for field in self.sections()["profile_fields"]]
        self.assertEqual(
            keys,
            ["profile:first_name", "profile:last_name", "profile:headline", "profile:phone",
             "profile:location", "profile:bio", "profile:linkedin_url"],
        )
        self.profile.user.first_name = "Already"
        self.profile.user.save()
        self.profile.headline = "Kept"
        self.profile.save()
        keys = [field["key"] for field in self.sections()["profile_fields"]]
        self.assertNotIn("profile:first_name", keys)
        self.assertNotIn("profile:headline", keys)
        self.assertIn("profile:last_name", keys)

    def test_skills_carry_a_key_a_level_a_category_and_a_duplicate_flag(self):
        category = SkillCategory.objects.get_or_create(name="Programming Languages", kind="technical")[0]
        UserSkill.objects.create(profile=self.profile, category=category, name="python")
        technical = self.sections()["technical_skills"]
        self.assertEqual(
            [(s["key"], s["name"], s["category"], s["level"], s["is_duplicate"]) for s in technical],
            [
                ("technical_skill:0", "Python", "Programming Languages", "Expert", True),
                ("technical_skill:1", "Zig", "Brand New Category", "Wizard", False),
                ("technical_skill:2", "Docker", "Other", "Intermediate", False),
            ],
        )
        soft = self.sections()["soft_skills"]
        self.assertEqual((soft[0]["key"], soft[0]["level"], soft[0]["is_duplicate"]), ("soft_skill:0", "Advanced", False))

    def test_the_same_name_in_the_other_kind_is_not_a_duplicate(self):
        soft = SkillCategory.objects.filter(kind="soft").first()
        UserSkill.objects.create(profile=self.profile, category=soft, name="Python")
        self.assertFalse(self.sections()["technical_skills"][0]["is_duplicate"])

    def test_languages_default_an_unknown_proficiency_and_flag_duplicates(self):
        UserLanguage.objects.create(
            profile=self.profile, language=Language.objects.create(name="french"),
            proficiency=UserLanguage.NATIVE,
        )
        languages = self.sections()["languages"]
        self.assertEqual(
            [(l["key"], l["name"], l["proficiency"], l["is_duplicate"]) for l in languages],
            [("language:0", "French", "Native", True), ("language:1", "Klingon", "Conversational", False)],
        )

    def test_incomplete_entries_are_dropped_and_repeats_are_flagged(self):
        WorkExperience.objects.create(
            profile=self.profile, job_title="backend engineer", company="ACME",
            start_date=date(2020, 1, 1),
        )
        Degree.objects.create(profile=self.profile, school="mcgill", degree="bsc")
        Certificate.objects.create(profile=self.profile, name="aws saa", issuing_organization="AMAZON")
        sections = self.sections()
        self.assertEqual(
            [(e["key"], e["is_duplicate"]) for e in sections["experience"]],
            [("experience:0", True), ("experience:1", False)],
        )
        self.assertEqual([(d["key"], d["is_duplicate"]) for d in sections["degrees"]], [("degree:0", True)])
        self.assertEqual(
            [(c["key"], c["is_duplicate"]) for c in sections["certificates"]], [("certificate:0", True)]
        )
        current = sections["experience"][0]
        self.assertEqual(current["end_date"], "")
        self.assertTrue(current["is_current"])
        self.assertEqual(current["highlights"], ["Built the payments API.", "Cut latency by half."])

    def test_an_import_with_nothing_in_it_yields_empty_sections(self):
        self.upload.ai_response = None
        empty = self.sections()
        self.assertTrue(all(value == [] for value in empty.values()))

    # --- the screen ---------------------------------------------------------
    def test_a_finished_import_shows_the_checklist(self):
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["resume_import"], self.upload)
        self.assertEqual(len(response.context["sections"]["technical_skills"]), 3)

    def test_an_import_still_running_shows_progress_instead(self):
        ResumeImport.objects.filter(pk=self.upload.pk).update(status=ResumeImport.STATUS_PROCESSING)
        task = AITask.start_for(self.profile, AITask.RESUME_IMPORT, self.upload, steps_total=2)
        response = self.client.get(self.url)
        self.assertIsNone(response.context["sections"])
        self.assertEqual(response.context["ai_task"], task)

    def test_applying_nothing_warns_and_keeps_the_page(self):
        response = self.client.post(self.url, {})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(flash(response), ["Select at least one item to add it to your profile."])
        self.assertIsNotNone(response.context["sections"])
        self.upload.refresh_from_db()
        self.assertEqual(self.upload.status, ResumeImport.STATUS_COMPLETED)

    def test_an_import_that_is_not_ready_cannot_be_applied(self):
        ResumeImport.objects.filter(pk=self.upload.pk).update(status=ResumeImport.STATUS_PENDING)
        response = self.client.post(self.url, {"selected": ["profile:headline"]})
        self.assertRedirects(response, self.url, fetch_redirect_response=False)
        self.assertEqual(Profile.objects.get(pk=self.profile.pk).headline, "")

    def test_applying_creates_what_was_ticked_and_says_so(self):
        response = self.client.post(
            self.url,
            {"selected": ["profile:first_name", "profile:headline", "technical_skill:0",
                          "technical_skill:1", "language:0", "experience:0", "degree:0",
                          "certificate:0"]},
        )
        self.assertRedirects(response, reverse("accounts:profile"), fetch_redirect_response=False)
        self.assertEqual(
            flash(response),
            ["Updated your profile, 2 skill(s), 1 language(s), 1 work experience entry, "
             "1 degree(s), 1 certificate(s) from your resume."],
        )
        self.upload.refresh_from_db()
        self.assertEqual(self.upload.status, ResumeImport.STATUS_APPLIED)
        self.assertIsNotNone(self.upload.applied_at)
        self.assertEqual(Profile.objects.get(pk=self.profile.pk).headline, "Backend engineer")
        self.assertEqual(User.objects.get(pk=self.user.pk).first_name, "Jane")

    def test_two_experiences_are_pluralised(self):
        response = self.client.post(self.url, {"selected": ["experience:0", "experience:1"]})
        self.assertEqual(flash(response), ["Updated 2 work experience entries from your resume."])

    def test_ticking_only_something_already_there_says_nothing_new(self):
        category = SkillCategory.objects.get_or_create(name="Programming Languages", kind="technical")[0]
        UserSkill.objects.create(profile=self.profile, category=category, name="Python")
        response = self.client.post(self.url, {"selected": ["technical_skill:0"]})
        self.assertEqual(flash(response), ["Updated nothing new from your resume."])

    def test_an_import_from_another_profile_is_a_404(self):
        _other_user, other = make_user_profile("cv-other@example.com")
        foreign = ResumeImport.objects.create(
            profile=other, file="resumes/y.txt", status=ResumeImport.STATUS_COMPLETED,
            ai_response=RESUME_REPLY,
        )
        url = reverse("resume:review", args=[foreign.pk])
        self.assertEqual(self.client.get(url).status_code, 404)
        self.assertEqual(self.client.post(url, {"selected": ["profile:headline"]}).status_code, 404)


class ApplySelectedTests(TestCase):
    """UC-04.4: the transactional import of exactly what the user ticked."""

    def setUp(self):
        self.user, self.profile = make_user_profile("apply@example.com", first_name="", last_name="")
        self.upload = ResumeImport.objects.create(
            profile=self.profile, file="resumes/x.txt", status=ResumeImport.STATUS_COMPLETED,
            ai_response=json.loads(json.dumps(RESUME_REPLY)),
        )

    def apply(self, *keys):
        from .services.importer import apply_selected

        return apply_selected(self.upload, self.profile, set(keys))

    def test_nothing_ticked_creates_nothing_but_still_marks_it_applied(self):
        counts = self.apply()
        self.assertEqual(
            counts, {"profile": False, "skills": 0, "languages": 0, "experience": 0,
                     "degrees": 0, "certificates": 0},
        )
        self.upload.refresh_from_db()
        self.assertEqual(self.upload.status, ResumeImport.STATUS_APPLIED)
        self.assertFalse(UserSkill.objects.exists())

    def test_names_go_to_the_account_and_the_rest_to_the_profile(self):
        counts = self.apply("profile:first_name", "profile:last_name", "profile:phone", "profile:bio")
        self.assertTrue(counts["profile"])
        self.user.refresh_from_db()
        self.profile.refresh_from_db()
        self.assertEqual((self.user.first_name, self.user.last_name), ("Jane", "Doe"))
        self.assertEqual((self.profile.phone, self.profile.bio), ("+1 555 0100", "Ships payment systems."))
        self.assertEqual(self.profile.headline, "", "only what was ticked")

    def test_an_empty_value_never_overwrites(self):
        self.user.first_name = "Keep"
        self.user.save()
        self.upload.ai_response["profile"]["first_name"] = ""
        self.apply("profile:first_name")
        self.user.refresh_from_db()
        self.assertEqual(self.user.first_name, "Keep")

    def test_skills_reuse_or_create_their_category_and_map_the_level(self):
        counts = self.apply("technical_skill:0", "technical_skill:1", "technical_skill:2", "soft_skill:0")
        self.assertEqual(counts["skills"], 4)
        python = UserSkill.objects.get(profile=self.profile, name="Python")
        self.assertEqual((python.category.name, python.category.kind, python.level), ("Programming Languages", "technical", UserSkill.EXPERT))
        zig = UserSkill.objects.get(profile=self.profile, name="Zig")
        self.assertEqual((zig.category.name, zig.level), ("Brand New Category", UserSkill.INTERMEDIATE))
        self.assertEqual(UserSkill.objects.get(name="Docker").category.name, "Other")
        self.assertEqual(UserSkill.objects.get(name="Mentoring").category.kind, "soft")
        self.assertEqual(SkillCategory.objects.filter(name="Programming Languages").count(), 1)

    def test_applying_the_same_skill_twice_does_not_duplicate_it(self):
        self.apply("technical_skill:0")
        counts = self.apply("technical_skill:0")
        self.assertEqual(counts["skills"], 0)
        self.assertEqual(UserSkill.objects.filter(name="Python").count(), 1)

    def test_languages_share_the_language_table_and_default_the_proficiency(self):
        existing = Language.objects.create(name="french")
        counts = self.apply("language:0", "language:1")
        self.assertEqual(counts["languages"], 2)
        self.assertEqual(UserLanguage.objects.get(language=existing).proficiency, "native")
        self.assertEqual(UserLanguage.objects.get(language__name="Klingon").proficiency, "conversational")
        self.assertEqual(Language.objects.filter(name__iexact="french").count(), 1)

    def test_experience_is_created_with_dates_type_and_ordered_highlights(self):
        counts = self.apply("experience:0", "experience:1")
        self.assertEqual(counts["experience"], 2)
        current = WorkExperience.objects.get(company="Acme")
        self.assertEqual(
            (current.job_title, current.employment_type, current.is_current, current.start_date, current.end_date),
            ("Backend Engineer", "full_time", True, date(2020, 1, 15), None),
        )
        self.assertEqual(
            [(h.order, h.text) for h in current.highlights.all()],
            [(0, "Built the payments API."), (1, "Cut latency by half.")],
        )
        past = WorkExperience.objects.get(company="Globex")
        self.assertEqual((past.employment_type, past.start_date, past.end_date), ("", date(2017, 1, 1), date(2019, 12, 1)))

    def test_an_experience_without_a_start_date_starts_today(self):
        self.upload.ai_response["experience"][0]["start_date"] = None
        self.apply("experience:0")
        self.assertEqual(WorkExperience.objects.get(company="Acme").start_date, timezone.localdate())

    def test_degrees_and_certificates(self):
        counts = self.apply("degree:0", "certificate:0")
        self.assertEqual((counts["degrees"], counts["certificates"]), (1, 1))
        degree = Degree.objects.get(profile=self.profile)
        self.assertEqual(
            (degree.school, degree.degree, degree.field_of_study, degree.grade, degree.start_date, degree.end_date),
            ("McGill", "BSc", "CS", "First", date(2013, 9, 1), date(2017, 5, 1)),
        )
        certificate = Certificate.objects.get(profile=self.profile)
        self.assertEqual(
            (certificate.name, certificate.issuing_organization, certificate.issue_date, certificate.expiry_date, certificate.credential_id),
            ("AWS SAA", "Amazon", date(2023, 1, 15), date(2026, 1, 15), "ABC"),
        )

    def test_incomplete_entries_are_skipped_even_when_ticked(self):
        counts = self.apply("experience:2", "degree:1", "certificate:1")
        self.assertEqual((counts["experience"], counts["degrees"], counts["certificates"]), (0, 0, 0))

    def test_unknown_keys_are_ignored(self):
        counts = self.apply("skill:99", "nonsense", "technical_skill:99")
        self.assertEqual(counts["skills"], 0)

    def test_everything_lands_in_the_profile_the_import_belongs_to(self):
        _other_user, other = make_user_profile("apply-other@example.com")
        self.apply("technical_skill:0", "language:0", "experience:0", "degree:0", "certificate:0")
        self.assertFalse(UserSkill.objects.filter(profile=other).exists())
        self.assertFalse(WorkExperience.objects.filter(profile=other).exists())

    def test_it_is_all_or_nothing(self):
        from .services import importer

        with patch.object(importer, "_parse_date", side_effect=RuntimeError("boom")):
            with self.assertRaises(RuntimeError):
                self.apply("technical_skill:0", "experience:0")
        self.assertFalse(UserSkill.objects.exists())
        self.upload.refresh_from_db()
        self.assertEqual(self.upload.status, ResumeImport.STATUS_COMPLETED)


@override_settings(STORAGES=TEST_STORAGES)
class TailoredResumeFlowTests(TailoredResumeTestMixin, TestCase):
    """UC-07.1 to UC-07.4 through the real task, the real writer and the real PDF."""

    def setUp(self):
        super().setUp()
        runtime_settings.invalidate()
        self.addCleanup(runtime_settings.invalidate)
        self.plan = Plan.objects.create(
            slug="free", name="Free", price_cents=0, is_default=True, monthly_tailored_resumes=1
        )
        self.add_profile_content()
        self.client.force_login(self.user)
        self.generate_url = reverse("resume:tailored_generate", args=[self.job.pk])
        self.editor_url = reverse("resume:tailored", args=[self.job.pk])

    def resumes_used(self):
        record = UsageRecord.objects.filter(
            user=self.user, metric=UsageMetric.TAILORED_RESUME
        ).first()
        return record.count if record else 0

    def generate(self, reply=None):
        with fake_deepseek(reply or (lambda request: dict(AI_RESPONSE))) as calls:
            response = self.client.post(self.generate_url)
        return response, calls

    def test_generating_writes_the_draft_and_spends_a_resume(self):
        response, calls = self.generate()
        self.assertRedirects(response, self.editor_url, fetch_redirect_response=False)
        self.assertEqual(flash(response), ["Writing your tailored resume — this takes a moment."])
        tailored = TailoredResume.objects.get(job=self.job)
        self.assertEqual(tailored.state, TailoredResume.STATE_COMPLETED)
        self.assertFalse(tailored.edited_by_user)
        self.assertEqual(tailored.ai_model, "deepseek-test")
        self.assertIsNotNone(tailored.generated_at)
        self.assertTrue(tailored.markdown.startswith("**Jane Doe**"))
        task = AITask.latest_for(tailored, AITask.TAILORED_RESUME)
        self.assertEqual((task.state, task.steps_total), (AITask.DONE, 1))
        self.assertEqual(self.resumes_used(), 1)
        self.assertEqual(len(calls), 1)

    def test_the_writer_gets_the_job_the_whole_profile_and_the_template(self):
        _response, (call,) = self.generate()
        self.assertIn("expert resume writer", call.system)
        self.assertEqual(call.temperature, 0.3)
        clause, payload = call.user.split("\n\n", 1)
        self.assertIn("in English", clause)
        body = json.loads(payload)
        self.assertEqual(set(body), {"job", "candidate_profile", "resume_template"})
        self.assertEqual(body["job"]["company"], "Globex")
        self.assertEqual(body["candidate_profile"]["contact"]["email"], "jane@example.com")
        self.assertIn("## ", body["resume_template"])

    def test_a_job_that_is_not_analysed_yet_is_sent_back(self):
        JobPost.objects.filter(pk=self.job.pk).update(status=JobPost.STATUS_PENDING)
        response, calls = self.generate()
        self.assertRedirects(response, self.job.get_absolute_url(), fetch_redirect_response=False)
        self.assertEqual(flash(response), ["Analyze this job post first, then generate a resume for it."])
        self.assertEqual(calls, [])
        self.assertEqual(self.resumes_used(), 0)

    def test_an_exhausted_allowance_or_a_plan_without_the_feature_explains_and_starts_nothing(self):
        runtime_settings.set_value("enforce_quotas", True)
        self.generate()
        response, calls = self.generate()
        self.assertRedirects(response, self.job.get_absolute_url(), fetch_redirect_response=False)
        self.assertIn("You've used all 1 of this month's tailored resumes.", flash(response)[-1])
        self.assertEqual(calls, [])
        self.assertEqual(self.resumes_used(), 1)

        Plan.objects.filter(pk=self.plan.pk).update(monthly_tailored_resumes=0)
        TailoredResume.objects.all().delete()
        response, _calls = self.generate()
        self.assertEqual(flash(response)[-1], "Your plan doesn't include this feature. Upgrade to use it.")
        self.assertFalse(TailoredResume.objects.exists())

    def test_the_kill_switch_blocks_generation(self):
        runtime_settings.set_value("ai_features_enabled", False)
        response, calls = self.generate()
        self.assertIn("temporarily unavailable", flash(response)[0])
        self.assertEqual(calls, [])
        self.assertFalse(TailoredResume.objects.exists())

    def test_a_model_failure_leaves_a_failed_draft_with_the_reason(self):
        _response, _calls = self.generate(lambda request: 401)
        tailored = TailoredResume.objects.get(job=self.job)
        self.assertEqual(tailored.state, TailoredResume.STATE_FAILED)
        self.assertIn("rejected our API key", tailored.error_message)
        self.assertEqual(AITask.latest_for(tailored, AITask.TAILORED_RESUME).state, AITask.FAILED)

    def test_a_profile_with_nothing_in_it_is_told_to_add_something_first(self):
        WorkExperience.objects.all().delete()
        UserSkill.objects.all().delete()
        Profile.objects.filter(pk=self.profile.pk).update(headline="", bio="")
        _response, calls = self.generate()
        self.assertEqual(calls, [])
        tailored = TailoredResume.objects.get(job=self.job)
        self.assertEqual(tailored.state, TailoredResume.STATE_FAILED)
        self.assertEqual(tailored.error_message, "Add some experience or skills to this profile first.")

    def test_regenerating_replaces_an_edited_draft(self):
        self.generate()
        TailoredResume.objects.filter(job=self.job).update(markdown="mine", edited_by_user=True)
        self.generate()
        tailored = TailoredResume.objects.get(job=self.job)
        self.assertFalse(tailored.edited_by_user)
        self.assertNotEqual(tailored.markdown, "mine")

    def test_the_editor_shows_the_draft(self):
        self.generate()
        response = self.client.get(self.editor_url)
        self.assertEqual(response.status_code, 200)
        context = response.context
        self.assertEqual(context["job"], self.job)
        self.assertEqual(context["tailored_resume"].job, self.job)
        self.assertEqual(context["form"].instance, context["tailored_resume"])
        self.assertEqual(context["ai_task"].kind, AITask.TAILORED_RESUME)

    def test_saving_text_that_did_not_change_is_not_an_edit(self):
        self.generate()
        saved = "**Jane Doe**\n\n## **SKILLS**\n\n- Python"
        TailoredResume.objects.filter(job=self.job).update(markdown=saved)
        response = self.client.post(self.editor_url, {"markdown": saved, "action": "save"})
        self.assertRedirects(response, self.editor_url, fetch_redirect_response=False)
        self.assertEqual(flash(response)[-1], "Saved your changes.")
        self.assertFalse(TailoredResume.objects.get(job=self.job).edited_by_user)

    def test_saving_different_text_marks_the_draft_as_edited(self):
        self.generate()
        saved = "**Jane Doe**\n\n## **SKILLS**\n\n- Python"
        TailoredResume.objects.filter(job=self.job).update(markdown=saved)
        self.client.post(self.editor_url, {"markdown": saved + "\n- Rust", "action": "save"})
        tailored = TailoredResume.objects.get(job=self.job)
        self.assertTrue(tailored.edited_by_user)
        self.assertTrue(tailored.markdown.endswith("- Rust"))

    def test_the_pdf_is_titled_after_the_job_and_authored_by_the_candidate(self):
        self.generate()
        response = self.client.get(reverse("resume:tailored_pdf", args=[self.job.pk]))
        from pypdf import PdfReader

        info = PdfReader(io.BytesIO(response.content)).metadata
        self.assertEqual((info.title, info.author), ("Senior Backend Engineer", "Jane Doe"))

    def test_the_pdf_title_falls_back_when_the_job_has_none(self):
        self.generate()
        JobPost.objects.filter(pk=self.job.pk).update(title="")
        response = self.client.get(reverse("resume:tailored_pdf", args=[self.job.pk]))
        from pypdf import PdfReader

        self.assertEqual(PdfReader(io.BytesIO(response.content)).metadata.title, "Resume")

    def test_deleting_the_resume_keeps_the_job(self):
        self.generate()
        response = self.client.post(reverse("resume:tailored_delete", args=[self.job.pk]))
        self.assertRedirects(response, self.job.get_absolute_url(), fetch_redirect_response=False)
        self.assertEqual(flash(response)[-1], "Deleted the tailored resume for this job.")
        self.assertFalse(TailoredResume.objects.exists())
        self.assertTrue(JobPost.objects.filter(pk=self.job.pk).exists())

    def test_the_list_shows_only_this_profiles_resumes(self):
        self.generate()
        other_user, other = make_user_profile("list-other@example.com")
        other_job = JobPost.objects.create(profile=other, title="Theirs", status=JobPost.STATUS_COMPLETED)
        TailoredResume.objects.create(profile=other, job=other_job, markdown="x")
        response = self.client.get(reverse("resume:tailored_list"))
        self.assertEqual([r.job for r in response.context["tailored_resumes"]], [self.job])
