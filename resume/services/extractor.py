"""Extracts plain text from an uploaded resume file (PDF, DOCX, or TXT)."""

import docx
from pypdf import PdfReader


class ResumeExtractError(Exception):
    """Raised for any user-facing failure reading an uploaded resume file."""


def extract_resume_text(django_file) -> str:
    name = (django_file.name or "").lower()
    django_file.seek(0)
    if name.endswith(".pdf"):
        return _extract_pdf(django_file)
    if name.endswith(".docx"):
        return _extract_docx(django_file)
    if name.endswith(".txt"):
        return _extract_txt(django_file)
    raise ResumeExtractError("Unsupported file type. Please upload a PDF, DOCX, or TXT file.")


def _extract_pdf(f) -> str:
    try:
        reader = PdfReader(f)
    except Exception:
        raise ResumeExtractError("Couldn't read that PDF — it may be corrupted.")

    if reader.is_encrypted:
        raise ResumeExtractError("That PDF is password-protected. Please upload an unlocked file.")

    pages_text = []
    for page in reader.pages:
        try:
            pages_text.append(page.extract_text() or "")
        except Exception:
            continue

    text = "\n".join(pages_text).strip()
    if len(text) < 50:
        raise ResumeExtractError(
            "We couldn't extract readable text from that PDF — it may be a scanned "
            "image. Try a text-based PDF, or upload a DOCX/TXT file instead."
        )
    return text


def _extract_docx(f) -> str:
    try:
        document = docx.Document(f)
    except Exception:
        raise ResumeExtractError("Couldn't read that DOCX file — it may be corrupted.")

    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text.strip())

    text = "\n".join(parts).strip()
    if len(text) < 50:
        raise ResumeExtractError("We couldn't find enough readable text in that document.")
    return text


def _extract_txt(f) -> str:
    raw = f.read()
    text = raw.decode("utf-8", errors="replace") if isinstance(raw, bytes) else raw
    text = text.strip()
    if len(text) < 50:
        raise ResumeExtractError("That file doesn't seem to contain enough text to analyze.")
    return text
